# -*- coding: utf-8 -*-
import os
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.abspath(__file__))
ASSET = os.path.join(ROOT, "報告素材")
PHOTO = os.path.join(ROOT, "成果照片")
OUT = os.path.join(ROOT, "0810物聯網成果報告.docx")
os.makedirs(ASSET, exist_ok=True)

FONT = r"C:\Windows\Fonts\msjh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msjhbd.ttc"

def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size, index=0)

def centered(draw, box, text, f, fill=(20, 35, 55)):
    x1, y1, x2, y2 = box
    bb = draw.textbbox((0, 0), text, font=f)
    draw.text(((x1 + x2 - (bb[2]-bb[0]))/2, (y1 + y2 - (bb[3]-bb[1]))/2), text, font=f, fill=fill)

def box(draw, xy, text, fill, outline=(40, 60, 80), f=None):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    centered(draw, xy, text, f or font(28, True))

def arrow(draw, start, end, fill=(50, 75, 100), width=5):
    draw.line([start, end], fill=fill, width=width)
    import math
    ang = math.atan2(end[1]-start[1], end[0]-start[0])
    a = 13
    p1 = (end[0] - a*math.cos(ang-0.5), end[1] - a*math.sin(ang-0.5))
    p2 = (end[0] - a*math.cos(ang+0.5), end[1] - a*math.sin(ang+0.5))
    draw.polygon([end, p1, p2], fill=fill)

def make_architecture(path):
    im = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 35), "ESP32 IoT Environment Monitor - Architecture", font=font(38, True), fill=(20, 45, 75))
    box(d, (70, 190, 390, 330), "DHT11\nGPIO14", (220, 242, 255), f=font(28, True))
    box(d, (70, 500, 390, 640), "Light Sensor\nGPIO33", (255, 244, 200), f=font(28, True))
    box(d, (610, 310, 990, 500), "ESP32 Dev Module\nRead / Display / Alert", (204, 236, 215), f=font(30, True))
    box(d, (1200, 175, 1515, 300), "OLED 0.96\nI2C: 21 / 22", (232, 224, 255), f=font(28, True))
    box(d, (1200, 400, 1515, 525), "Wi-Fi AP\nSSID: a", (222, 243, 247), f=font(28, True))
    box(d, (1200, 625, 1515, 750), "LINE Messaging API\nPush notification", (255, 224, 232), f=font(27, True))
    box(d, (610, 650, 990, 790), "LINE User\nAlert received", (255, 235, 210), f=font(30, True))
    arrow(d, (390, 260), (610, 370))
    arrow(d, (390, 570), (610, 440))
    arrow(d, (990, 365), (1200, 240))
    arrow(d, (990, 405), (1200, 460))
    arrow(d, (1355, 525), (1355, 625))
    arrow(d, (1200, 690), (990, 720))
    d.text((1035, 330), "I2C display", font=font(23), fill=(65, 75, 90))
    d.text((1010, 545), "HTTPS", font=font(23), fill=(65, 75, 90))
    im.save(path)

def make_flow(path):
    im = Image.new("RGB", (1500, 1450), "white")
    d = ImageDraw.Draw(im)
    d.text((50, 30), "Runtime Flow", font=font(42, True), fill=(20, 45, 75))
    boxes = [
        ((470, 110, 1030, 220), "Start / connect Wi-Fi", (222, 243, 247)),
        ((470, 285, 1030, 395), "Read DHT11 + light sensor", (220, 242, 255)),
        ((470, 460, 1030, 570), "Update OLED values", (232, 224, 255)),
        ((470, 635, 1030, 745), "Temperature > 28 OR humidity > 70?", (255, 244, 200)),
        ((120, 850, 620, 970), "No: continue monitoring", (226, 242, 226)),
        ((880, 850, 1380, 970), "Yes: abnormal state", (255, 224, 232)),
        ((880, 1050, 1380, 1170), "30 seconds elapsed?", (255, 244, 200)),
        ((880, 1250, 1380, 1370), "Show LINE sending\nPush Chinese alert", (255, 235, 210)),
    ]
    for xy, txt, col in boxes:
        box(d, xy, txt, col, f=font(28, True))
    arrow(d, (750, 220), (750, 285)); arrow(d, (750, 395), (750, 460)); arrow(d, (750, 570), (750, 635))
    arrow(d, (470, 690), (370, 850)); arrow(d, (1030, 690), (1130, 850))
    arrow(d, (1130, 970), (1130, 1050)); arrow(d, (1130, 1170), (1130, 1250))
    arrow(d, (880, 1110), (620, 910))
    d.text((430, 755), "No", font=font(25, True), fill=(40, 90, 50))
    d.text((1050, 755), "Yes", font=font(25, True), fill=(150, 45, 65))
    d.text((690, 1160), "No: wait", font=font(24, True), fill=(70, 70, 70))
    im.save(path)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_font(run, size=12, bold=False, color=None):
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_p(doc, text, size=12, bold=False, align=None, color=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, 17 if level == 1 else 14, True, (31, 78, 121))
    return p

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 10, False, (90, 90, 90))

def add_picture(doc, path, width_cm):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))

def add_table(doc, rows, widths=None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, val in enumerate(rows[0]):
        hdr[i].text = ''
        set_cell_shading(hdr[i], '1F4E79')
        r = hdr[i].paragraphs[0].add_run(val)
        set_font(r, 11, True, (255,255,255))
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(val)
            set_font(r, 11)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table

ARCHITECTURE_IMAGE = os.path.join(ASSET, "架構圖_AI_中文.png")
FLOW_IMAGE = os.path.join(ASSET, "流程圖_AI_中文.png")

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)

# Cover
for _ in range(4): doc.add_paragraph()
add_p(doc, "0810 物聯網成果報告", 28, True, WD_ALIGN_PARAGRAPH.CENTER, (31,78,121), 18)
add_p(doc, "ESP32 環境感測、OLED 顯示與 LINE 異常通知系統", 17, False, WD_ALIGN_PARAGRAPH.CENTER, (70,70,70), 18)
add_p(doc, "專案最終版本：25_dht_line", 14, False, WD_ALIGN_PARAGRAPH.CENTER, (90,90,90), 12)
add_p(doc, "報告日期：2026 年 8 月 10 日", 13, False, WD_ALIGN_PARAGRAPH.CENTER, (90,90,90), 12)
doc.add_page_break()

add_heading(doc, "摘要", 1)
add_p(doc, "本專案以 ESP32 Dev Module 為核心，整合 DHT11 溫濕度感測器、光敏電阻與 0.96 吋 OLED 顯示器，建立可即時監測環境狀態的物聯網系統。系統每秒讀取感測值並更新 OLED；當溫度超過 28°C 或濕度超過 70% 時，透過 Wi‑Fi 呼叫 LINE Messaging API，向指定使用者傳送中文警告訊息，異常持續時每 30 秒最多通知一次。專案亦完成 ThingSpeak 與 Google Sheet 雲端記錄版本，形成由本地顯示、雲端記錄到即時告警的完整實作歷程。")

add_heading(doc, "一、專案目標", 1)
for t in [
    "建立 ESP32 基礎感測與輸出控制能力，熟悉 GPIO、ADC、I²C 與 Wi‑Fi。",
    "將 DHT11 溫度、濕度及光敏電阻亮度轉換為可讀數值，並在 OLED 上呈現。",
    "完成雲端資料記錄：ThingSpeak 圖表與 Google Sheet 表格。",
    "加入異常判斷與 LINE 即時通知，提高系統的實用性。",
]:
    add_p(doc, "• " + t)

add_heading(doc, "二、專案發展歷程", 1)
add_table(doc, [
    ("階段", "內容"),
    ("01–08 基礎硬體", "Hello、LED/RGB LED、夜間控制、PIR、DHT11、I²C、OLED。"),
    ("09–17 整合控制", "DHT11＋OLED、LED 警報、PWM、蜂鳴器、超音波與伺服相關練習。"),
    ("18–21 網路應用", "Bluetooth 指令、PM2.5、OLED 空氣品質與氣象資料。"),
    ("22–24 雲端記錄", "環境資料轉換亮度百分比，依序完成 ThingSpeak 與 Google Sheet 上傳。"),
    ("25 最終成果", "異常條件判斷、中文 LINE 通知、OLED 狀態回饋與 30 秒通知間隔。"),
])

add_heading(doc, "三、硬體與軟體配置", 1)
add_table(doc, [
    ("項目", "設定／用途"),
    ("控制器", "ESP32 Dev Module"),
    ("DHT11", "GPIO14；讀取溫度與濕度"),
    ("光敏電阻", "GPIO33 ADC；0–4095 線性轉換為 0–100% 亮度"),
    ("OLED", "0.96 吋 SSD1306 128×64；SDA GPIO21、SCL GPIO22"),
    ("通訊", "Wi‑Fi SSID a；LINE Messaging API HTTPS Push"),
    ("函式庫", "SimpleDHT、U8g2、WiFi、WiFiClientSecure"),
])

add_heading(doc, "四、系統架構", 1)
add_p(doc, "ESP32 讀取兩種感測器資料後，一方面透過 I²C 更新 OLED，另一方面經由 Wi‑Fi 連線至 LINE Messaging API。當異常條件成立，系統將警告訊息推送給指定使用者。")
add_picture(doc, ARCHITECTURE_IMAGE, 16.2)
add_caption(doc, "圖 1　ESP32 環境監測與 LINE 通知系統架構")

add_heading(doc, "五、程式流程", 1)
add_p(doc, "系統啟動後先連線 Wi‑Fi，OLED 顯示連線進度；連線完成後每秒讀取感測器並更新畫面。每次讀取都檢查溫度與濕度是否超過門檻，只有在異常且距離前次通知已達 30 秒時才傳送 LINE，避免短時間大量重複通知。")
add_picture(doc, FLOW_IMAGE, 16.2)
add_caption(doc, "圖 2　程式執行流程圖")

add_heading(doc, "六、主要功能說明", 1)
for t in [
    "感測與顯示：DHT11 每秒讀取溫度與濕度；光敏電阻 ADC 值轉換成 0–100% 亮度。OLED 以三欄方式顯示溫度計、水滴與太陽圖示及數值。",
    "異常判斷：溫度 > 28°C 或濕度 > 70% 即視為異常；DHT11 讀取失敗時不觸發通知。",
    "LINE 通知：透過 api.line.me 的 Push API 傳送中文訊息「警告！溫度：XX C，濕度：XX %」給指定 User ID。",
    "通知節流：異常持續時每 30 秒最多傳送一次；OLED 底部顯示 Monitoring、LINE sending、LINE sent 或 LINE error。",
    "網路處理：啟動時顯示 Wi‑Fi 連線過程；若異常時 Wi‑Fi 中斷，程式會先嘗試重新連線。",
]:
    add_p(doc, "• " + t)

add_heading(doc, "七、成果展示", 1)
add_p(doc, "以下照片為最終版本實際測試成果。LINE 對話畫面可看到中文異常通知，證明系統已完成感測、判斷、網路傳送與使用者端接收。")
add_picture(doc, os.path.join(PHOTO, "25_dht_line.png"), 10.5)
add_caption(doc, "圖 3　LINE 中文異常通知成果照片")

add_p(doc, "前期雲端記錄版本亦完成驗證：Google Sheet 可逐筆記錄時間、溫度、濕度與亮度；ThingSpeak 可用圖表觀察三項資料的變化。")
add_picture(doc, os.path.join(PHOTO, "24_google_dht_light_oled.png"), 10.5)
add_caption(doc, "圖 4　Google Sheet 環境資料記錄")
add_picture(doc, os.path.join(PHOTO, "23_dht_light_oled_thingspeak.png"), 16.2)
add_caption(doc, "圖 5　ThingSpeak 三欄資料圖表")

add_heading(doc, "八、測試結果", 1)
add_table(doc, [
    ("測試項目", "結果"),
    ("DHT11 溫度／濕度讀取", "正常，OLED 每秒更新。"),
    ("光敏電阻亮度換算", "正常，ADC 值可轉換為 0–100%。"),
    ("OLED 顯示", "正常，三欄圖示與數值可同時呈現。"),
    ("Wi‑Fi 連線", "正常，啟動時顯示連線過程。"),
    ("異常 LINE 通知", "溫度 29°C、濕度 75% 測試成功收到中文通知。"),
    ("通知間隔", "異常持續時以 30 秒為最短通知間隔。"),
])

add_heading(doc, "九、問題與改進方向", 1)
for t in [
    "目前 LINE Channel Access Token 直接寫在韌體中，正式部署時應改採安全的設定管理方式，並定期更新 Token。",
    "亮度百分比是以 ADC 線性換算，若光敏電阻分壓方向不同，需將百分比反向或重新校正。",
    "可增加異常恢復通知、LINE 貼圖或更多感測欄位；也可加入 Google Sheet 與 LINE 的整合記錄。",
    "可加入時間戳記、資料校正與離線緩存，提升長時間運作的可靠度。",
]:
    add_p(doc, "• " + t)

add_heading(doc, "十、結論", 1)
add_p(doc, "本專案完成 ESP32 從基礎 GPIO 與感測器操作，到 OLED 圖形化顯示、雲端資料記錄及 LINE 即時告警的完整整合。最終版本能在環境條件超過設定門檻時主動通知使用者，並以 30 秒間隔控制通知頻率，兼顧即時性與避免訊息過量。透過本專案，已具備物聯網系統中感測、處理、顯示、通訊與通知等核心功能的實作經驗。")

add_p(doc, "主要程式：25_dht_line/25_dht_line.ino", 10, False, color=(100,100,100), space_after=2)
add_p(doc, "報告檔案編碼：UTF-8（中文內容以 Unicode 儲存於 Word 文件）", 10, False, color=(100,100,100), space_after=2)

# Footer
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("0810 物聯網成果報告　|　ESP32 + DHT11 + OLED + LINE")
    set_font(r, 9, False, (120,120,120))

doc.save(OUT)
print(OUT)
